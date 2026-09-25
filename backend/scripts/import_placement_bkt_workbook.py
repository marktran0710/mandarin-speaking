"""Import synthetic placement responses into the placement/BKT ledger.

Validation is the default. Database writes require ``--apply`` and are
performed as one transaction after the complete workbook and current active
placement blueprint have been checked.

Examples::

    python -m scripts.import_placement_bkt_workbook \
        "D:/data/responses_to_import_only_40_students_28q_placementtest_bkt.xlsx"
    python -m scripts.import_placement_bkt_workbook \
        "D:/data/responses_to_import_only_40_students_28q_placementtest_bkt.xlsx" \
        --apply
    # After the active blueprint changed: drop this importer's earlier data
    # for the workbook's students and import again, in one transaction.
    python -m scripts.import_placement_bkt_workbook \
        "D:/data/responses_to_import_only_40_students_28q_placementtest_bkt.xlsx" \
        --apply --replace
"""

from __future__ import annotations

import argparse
import copy
import math
import os
import secrets
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping

from dotenv import load_dotenv
import openpyxl
from psycopg.types.json import Jsonb


BACKEND_DIR = Path(__file__).resolve().parents[1]
load_dotenv(BACKEND_DIR / ".env")
sys.path.insert(0, str(BACKEND_DIR))

import security.auth as auth  # noqa: E402
from analytics.learner_model.bkt.mastery import (  # noqa: E402
    rebuild_student_vocabulary_mastery,
    upsert_raw_responses,
)
from db import connect_db  # noqa: E402
from domain.vocabulary.assessment import normalize_answer  # noqa: E402


IMPORT_RESOLVER_VERSION = "placement-workbook-import-v1"
EXPECTED_HEADERS = (
    "student_id",
    "student_name",
    "placement_session_id",
    "source_story_id",
    "item_id",
    "mode",
    "selected_answer",
    "answered_at",
    "time_ms",
)
EXPECTED_STUDENT_IDS = tuple(f"SIM{index:03d}" for index in range(1, 41))
EXPECTED_QUESTION_COUNT = 28
SUPPORTED_TYPES = {
    "basic_meaning_mcq": (1, "meaning"),
    "character_to_pinyin_typing": (2, "pinyin_production"),
    "context_cloze_mcq": (3, "contextual_recall"),
}


class ImportConflictError(ValueError):
    """Raised when an import would silently overwrite immutable learner data."""


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _looks_local(url: str) -> bool:
    return "@127.0.0.1" in url or "@localhost" in url or "@::1" in url


def _as_time_ms(value: object, row_number: int) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"row {row_number}: time_ms must be a non-negative integer")
    if not math.isfinite(float(value)) or int(value) != value or value < 0:
        raise ValueError(f"row {row_number}: time_ms must be a non-negative integer")
    return int(value)


def read_workbook(path: Path) -> list[dict[str, object]]:
    """Read the exact response shape used by the import contract."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = tuple(_text(value) for value in next(iterator))
        except StopIteration as exc:
            raise ValueError("Workbook has no header row.") from exc
        if header != EXPECTED_HEADERS:
            raise ValueError(
                "Workbook headers must be exactly: " + ", ".join(EXPECTED_HEADERS)
            )

        rows: list[dict[str, object]] = []
        for row_number, values in enumerate(iterator, start=2):
            if not any(value is not None for value in values):
                continue
            if len(values) != len(EXPECTED_HEADERS):
                raise ValueError(f"row {row_number}: expected {len(EXPECTED_HEADERS)} columns")
            row = dict(zip(EXPECTED_HEADERS, values))
            for field in EXPECTED_HEADERS[:7]:
                row[field] = _text(row[field])
            row["time_ms"] = _as_time_ms(row["time_ms"], row_number)
            if not all(_text(row[field]) for field in EXPECTED_HEADERS[:7]):
                raise ValueError(f"row {row_number}: required response fields cannot be blank")
            rows.append(row)
        return rows
    finally:
        workbook.close()


def load_active_blueprint(db: Any) -> dict[str, Any]:
    row = db.execute(
        "SELECT id, revision, questions FROM placement_test_blueprints WHERE id = %s",
        ("active",),
    ).fetchone()
    if not row or not isinstance(row.get("questions"), list) or not row["questions"]:
        raise ValueError("The active placement blueprint is not configured.")
    questions = sorted(
        (copy.deepcopy(question) for question in row["questions"] if isinstance(question, dict)),
        key=lambda question: int(question.get("position", 0)),
    )
    if len(questions) != EXPECTED_QUESTION_COUNT:
        raise ValueError(
            f"Active placement blueprint has {len(questions)} questions; "
            f"expected {EXPECTED_QUESTION_COUNT}."
        )
    question_ids = [_text(question.get("questionId")) for question in questions]
    if not all(question_ids) or len(set(question_ids)) != len(question_ids):
        raise ValueError("Active placement blueprint must have unique questionId values.")
    return {"revision": int(row["revision"]), "questions": questions}


def _accepted_answers(question: Mapping[str, Any]) -> list[str]:
    accepted = question.get("acceptedAnswers") or []
    if isinstance(accepted, str):
        accepted = [accepted]
    return [str(value) for value in accepted] or [_text(question.get("correctAnswer"))]


def _grade(question: Mapping[str, Any], selected_answer: str) -> bool:
    selected = normalize_answer(selected_answer)
    return bool(selected) and any(
        selected == normalize_answer(answer) for answer in _accepted_answers(question)
    )


def _question_identity(question: Mapping[str, Any]) -> dict[str, Any]:
    return {
        key: question.get(key)
        for key in (
            "questionId",
            "sourceStoryId",
            "sourceWordId",
            "round",
            "tier",
            "questionType",
            "answerFormat",
            "correctAnswer",
            "acceptedAnswers",
            "options",
        )
    }


def _question_snapshot_matches(left: Any, right: Any) -> bool:
    if not isinstance(left, list) or not isinstance(right, list):
        return False
    left_questions = sorted(
        (question for question in left if isinstance(question, dict)),
        key=lambda question: int(question.get("position", 0)),
    )
    right_questions = sorted(
        (question for question in right if isinstance(question, dict)),
        key=lambda question: int(question.get("position", 0)),
    )
    return [_question_identity(question) for question in left_questions] == [
        _question_identity(question) for question in right_questions
    ]


def _response_snapshot_matches(left: Any, right: Any) -> bool:
    if not isinstance(left, list) or not isinstance(right, list):
        return False
    fields = ("questionId", "selectedAnswer", "timeMs", "correct", "position")
    return [
        {field: response.get(field) for field in fields}
        for response in left
        if isinstance(response, dict)
    ] == [
        {field: response.get(field) for field in fields}
        for response in right
        if isinstance(response, dict)
    ]


def _canonical_response_snapshot(
    question: Mapping[str, Any], selected_answer: str, time_ms: int, correct: bool
) -> dict[str, Any]:
    return {
        "questionId": _text(question["questionId"]),
        "selectedAnswer": selected_answer,
        "timeMs": time_ms,
        "answeredAt": None,
        "correct": correct,
        "position": int(question["position"]),
    }


def _response_row(
    student_id: str,
    session_id: str,
    question: Mapping[str, Any],
    selected_answer: str,
    time_ms: int,
    correct: bool,
    imported_at: str,
    occurred_at_utc: datetime,
) -> dict[str, Any]:
    question_type = _text(question.get("questionType"))
    round_number, dimension = SUPPORTED_TYPES[question_type]
    question_prompt = _text(question.get("prompt"))
    if question_type == "character_to_pinyin_typing" and not question_prompt:
        question_prompt = f"Type the pinyin for {_text(question.get('targetWord'))}."
    return {
        "student_id": student_id,
        "word_id": _text(question["sourceWordId"]),
        "word": _text(question.get("targetWord")) or _text(question["sourceWordId"]),
        "lesson_id": _text(question["sourceStoryId"]),
        "quiz_id": session_id,
        "attempt_id": session_id,
        "item_id": _text(question["questionId"]),
        "question_type": question_type,
        "selected_answer": selected_answer,
        "correct_answer": _text(question.get("correctAnswer")),
        "presented_options": [str(value) for value in (question.get("options") or [])],
        "question_prompt": question_prompt,
        "answered_at": imported_at,
        "bkt_eligible": True,
        "diagnostic_exposure_id": f"placement:{session_id}:{question['questionId']}",
        "bkt_eligibility_errors": [],
        "correct": correct,
        "response_time_ms": time_ms,
        "occurred_at": imported_at,
        "occurred_at_utc": occurred_at_utc,
        "evidence_origin": "synthetic",
        "resolver_version": IMPORT_RESOLVER_VERSION,
        "attempt_order": int(question["position"]) - 1,
        "quiz_level": _text(question["tier"]),
        "quiz_mode": _text(question["tier"]),
        "round_type": str(round_number),
        "knowledge_dimension": dimension,
        "activity_type": "diagnostic",
        "research_study_id": None,
    }


def build_import_plan(
    rows: Iterable[Mapping[str, object]],
    blueprint: Mapping[str, Any],
    *,
    imported_at: str,
    expected_student_ids: Iterable[str] = EXPECTED_STUDENT_IDS,
) -> dict[str, Any]:
    """Validate workbook facts and resolve them into immutable DB payloads."""
    rows = list(rows)
    questions = list(blueprint["questions"])
    by_item = {_text(question.get("questionId")): question for question in questions}
    expected_ids = set(expected_student_ids)
    by_student: dict[str, list[Mapping[str, object]]] = defaultdict(list)
    for row in rows:
        student_id = _text(row.get("student_id"))
        if student_id not in expected_ids:
            raise ValueError(f"Unexpected student_id: {student_id or '<empty>'}")
        by_student[student_id].append(row)

    if set(by_student) != expected_ids:
        missing = sorted(expected_ids - set(by_student))
        extra = sorted(set(by_student) - expected_ids)
        raise ValueError(f"Student coverage mismatch; missing={missing}, extra={extra}")

    attempts: list[dict[str, Any]] = []
    response_rows: list[dict[str, Any]] = []
    student_payloads: list[dict[str, str]] = []
    correct_by_mode: Counter[str] = Counter()
    source_aliases: Counter[tuple[str, str]] = Counter()
    occurred_at_utc = datetime.fromisoformat(imported_at.replace("Z", "+00:00")).astimezone(timezone.utc)

    for student_id in sorted(expected_ids):
        student_rows = by_student[student_id]
        if len(student_rows) != len(questions):
            raise ValueError(
                f"{student_id}: expected {len(questions)} rows, found {len(student_rows)}"
            )
        names = {_text(row.get("student_name")) for row in student_rows}
        sessions = {_text(row.get("placement_session_id")) for row in student_rows}
        if len(names) != 1:
            raise ValueError(f"{student_id}: student_name changes within the session")
        if len(sessions) != 1:
            raise ValueError(f"{student_id}: placement_session_id changes within the session")
        session_id = next(iter(sessions))
        if not session_id:
            raise ValueError(f"{student_id}: placement_session_id cannot be blank")
        by_item_row: dict[str, Mapping[str, object]] = {}
        for row in student_rows:
            item_id = _text(row.get("item_id"))
            if item_id in by_item_row:
                raise ValueError(f"{student_id}: duplicate item_id {item_id}")
            if item_id not in by_item:
                raise ValueError(f"{student_id}: item_id {item_id} is not in the active blueprint")
            by_item_row[item_id] = row

        missing_items = sorted(set(by_item) - set(by_item_row))
        if missing_items:
            raise ValueError(f"{student_id}: missing item_id values: {', '.join(missing_items)}")

        student_responses: list[dict[str, Any]] = []
        total_time_ms = 0
        for question in questions:
            item_id = _text(question["questionId"])
            row = by_item_row[item_id]
            mode = _text(row.get("mode")).casefold()
            expected_tier = _text(question.get("tier")).casefold()
            if mode != expected_tier:
                raise ValueError(
                    f"{student_id}/{item_id}: mode {mode!r} does not match active tier {expected_tier!r}"
                )
            question_type = _text(question.get("questionType"))
            if question_type not in SUPPORTED_TYPES:
                raise ValueError(f"{item_id}: unsupported questionType {question_type!r}")
            selected_answer = _text(row.get("selected_answer"))
            time_ms = _as_time_ms(row.get("time_ms"), 0)
            correct = _grade(question, selected_answer)
            source_aliases[(_text(row.get("source_story_id")), _text(question["sourceStoryId"]))] += 1
            correct_by_mode[mode] += int(correct)
            total_time_ms += time_ms
            student_responses.append(
                _canonical_response_snapshot(question, selected_answer, time_ms, correct)
            )
            response_rows.append(
                _response_row(
                    student_id,
                    session_id,
                    question,
                    selected_answer,
                    time_ms,
                    correct,
                    imported_at,
                    occurred_at_utc,
                )
            )

        student_payloads.append({"id": student_id, "name": next(iter(names))})
        attempts.append(
            {
                "id": session_id,
                "student_id": student_id,
                "blueprint_revision": int(blueprint["revision"]),
                "question_snapshot": copy.deepcopy(questions),
                "response_snapshot": [
                    {**response, "answeredAt": imported_at}
                    for response in student_responses
                ],
                "started_at": imported_at,
                "completed_at": imported_at,
                "total_questions": len(questions),
                "correct_count": sum(1 for response in student_responses if response["correct"]),
                "total_time_ms": total_time_ms,
            }
        )

    total_correct = sum(attempt["correct_count"] for attempt in attempts)
    return {
        "students": student_payloads,
        "attempts": attempts,
        "response_rows": response_rows,
        "summary": {
            "student_count": len(student_payloads),
            "question_count": len(questions),
            "response_count": len(response_rows),
            "correct_count": total_correct,
            "incorrect_count": len(response_rows) - total_correct,
            "correct_by_mode": dict(correct_by_mode),
            "source_aliases": dict(source_aliases),
        },
    }


def _student_conflict(existing: Mapping[str, Any], expected: Mapping[str, str]) -> str | None:
    if existing.get("name") != expected["name"]:
        return f"{expected['id']}: existing name is {existing.get('name')!r}, expected {expected['name']!r}"
    if not bool(existing.get("is_test_account")):
        return f"{expected['id']}: existing account is not marked is_test_account"
    if existing.get("status") != "active":
        return f"{expected['id']}: existing account status is {existing.get('status')!r}, expected 'active'"
    return None


def _existing_state(db: Any, plan: Mapping[str, Any]) -> dict[str, Any]:
    student_ids = [student["id"] for student in plan["students"]]
    session_ids = [attempt["id"] for attempt in plan["attempts"]]
    students = {
        row["id"]: row
        for row in db.execute(
            "SELECT id, name, status, is_test_account, password_reset_required "
            "FROM students WHERE id = ANY(%s)",
            (student_ids,),
        ).fetchall()
    }
    names = [student["name"].casefold() for student in plan["students"]]
    name_rows = db.execute(
        "SELECT id, name FROM students WHERE lower(name) = ANY(%s)",
        (names,),
    ).fetchall()
    for row in name_rows:
        if row["id"] not in student_ids:
            raise ImportConflictError(
                f"Student name {row['name']!r} already belongs to {row['id']}"
            )
    for expected in plan["students"]:
        existing = students.get(expected["id"])
        if existing:
            conflict = _student_conflict(existing, expected)
            if conflict:
                raise ImportConflictError(conflict)

    attempts = {
        row["id"]: row
        for row in db.execute(
            "SELECT id, student_id, blueprint_revision, question_snapshot, response_snapshot, "
            "status, total_questions, correct_count, total_time_ms "
            "FROM placement_test_attempts WHERE id = ANY(%s)",
            (session_ids,),
        ).fetchall()
    }
    response_rows = db.execute(
        "SELECT student_id, quiz_id, attempt_order, item_id, word_id, lesson_id, "
        "question_type, selected_answer, correct_answer, correct, response_time_ms, "
        "quiz_level, quiz_mode, evidence_origin, resolver_version "
        "FROM vocab_quiz_responses WHERE quiz_id = ANY(%s)",
        (session_ids,),
    ).fetchall()
    responses_by_session: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in response_rows:
        responses_by_session[row["quiz_id"]].append(row)

    skip_sessions: set[str] = set()
    plan_by_session = {attempt["id"]: attempt for attempt in plan["attempts"]}
    for session_id, attempt in plan_by_session.items():
        existing_attempt = attempts.get(session_id)
        existing_responses = responses_by_session.get(session_id, [])
        if existing_attempt is None:
            if existing_responses:
                raise ImportConflictError(
                    f"{session_id}: response ledger rows exist without a placement attempt"
                )
            continue
        if existing_attempt["student_id"] != attempt["student_id"]:
            raise ImportConflictError(
                f"{session_id}: existing placement attempt belongs to another student"
            )
        if existing_attempt["status"] != "completed":
            raise ImportConflictError(f"{session_id}: existing placement attempt is not completed")
        if not _question_snapshot_matches(
            existing_attempt["question_snapshot"], attempt["question_snapshot"]
        ):
            raise ImportConflictError(f"{session_id}: question snapshot differs from the active blueprint")
        if not _response_snapshot_matches(
            existing_attempt["response_snapshot"], attempt["response_snapshot"]
        ):
            raise ImportConflictError(f"{session_id}: response snapshot differs from the workbook")
        if len(existing_responses) != len(attempt["response_snapshot"]):
            raise ImportConflictError(f"{session_id}: response ledger row count differs from the workbook")
        expected_rows = {
            (row["student_id"], row["quiz_id"], row["attempt_order"]): row
            for row in plan["response_rows"]
            if row["quiz_id"] == session_id
        }
        actual_rows = {
            (row["student_id"], row["quiz_id"], row["attempt_order"]): row
            for row in existing_responses
        }
        if set(expected_rows) != set(actual_rows):
            raise ImportConflictError(f"{session_id}: response slots differ from the workbook")
        for slot, expected in expected_rows.items():
            actual = actual_rows[slot]
            for field in (
                "item_id",
                "word_id",
                "lesson_id",
                "question_type",
                "selected_answer",
                "correct_answer",
                "correct",
                "response_time_ms",
                "quiz_level",
                "quiz_mode",
            ):
                if actual.get(field) != expected[field]:
                    raise ImportConflictError(f"{session_id}: immutable response field {field} differs")
            if actual.get("evidence_origin") != "synthetic" or actual.get("resolver_version") != IMPORT_RESOLVER_VERSION:
                raise ImportConflictError(f"{session_id}: response provenance differs from this importer")
        skip_sessions.add(session_id)
    return {"students": students, "attempts": attempts, "skip_sessions": skip_sessions}


_REPLACEABLE_ATTEMPTS = """
    SELECT a.id FROM placement_test_attempts AS a
    WHERE a.student_id = ANY(%(students)s)
      AND (
        a.id = ANY(%(sessions)s)
        OR EXISTS (
            SELECT 1 FROM vocab_quiz_responses AS r
            WHERE r.attempt_id = a.id AND r.evidence_origin = 'synthetic' AND r.resolver_version = %(resolver)s
        )
      )
      AND NOT EXISTS (
        SELECT 1 FROM vocab_quiz_responses AS r
        WHERE r.attempt_id = a.id
          AND NOT (r.evidence_origin = 'synthetic' AND r.resolver_version = %(resolver)s)
      )
"""


def _replace_params(plan: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "students": [student["id"] for student in plan["students"]],
        "sessions": [attempt["id"] for attempt in plan["attempts"]],
        "resolver": IMPORT_RESOLVER_VERSION,
    }


def previous_import_counts(db: Any, plan: Mapping[str, Any]) -> dict[str, int]:
    """Count what --replace would delete (see delete_previous_import)."""
    params = _replace_params(plan)
    attempts = db.execute(f"SELECT count(*) AS n FROM ({_REPLACEABLE_ATTEMPTS}) AS a", params).fetchone()
    responses = db.execute(
        "SELECT count(*) AS n FROM vocab_quiz_responses WHERE student_id = ANY(%(students)s) "
        "AND evidence_origin = 'synthetic' AND resolver_version = %(resolver)s",
        params,
    ).fetchone()
    return {"attempts": int(attempts["n"]), "responses": int(responses["n"])}


def delete_previous_import(db: Any, plan: Mapping[str, Any]) -> dict[str, int]:
    """Remove this importer's earlier data for the workbook's students.

    Responses are deleted only when they are synthetic rows written by this
    importer. An attempt is deleted only when it belongs to a workbook student,
    is one of the workbook's sessions or held this importer's rows, and holds
    no other evidence - so real learner evidence is never touched. Student
    accounts are kept so their ids stay stable.
    """
    params = _replace_params(plan)
    attempt_ids = [row["id"] for row in db.execute(_REPLACEABLE_ATTEMPTS, params).fetchall()]
    responses = db.execute(
        "DELETE FROM vocab_quiz_responses WHERE student_id = ANY(%(students)s) "
        "AND evidence_origin = 'synthetic' AND resolver_version = %(resolver)s",
        params,
    ).rowcount
    attempts = db.execute(
        "DELETE FROM placement_test_attempts WHERE id = ANY(%s)", (attempt_ids,)
    ).rowcount
    return {"attempts": attempts, "responses": responses}


def _insert_student(db: Any, student: Mapping[str, str]) -> None:
    db.execute(
        "INSERT INTO students "
        "(id, name, password, password_reset_required, status, is_test_account) "
        "VALUES (%s, %s, %s, true, 'active', true)",
        (student["id"], student["name"], auth.hash_password(secrets.token_urlsafe(32))),
    )


def _insert_attempt(db: Any, attempt: Mapping[str, Any]) -> None:
    db.execute(
        """
        INSERT INTO placement_test_attempts
            (id, student_id, blueprint_revision, question_snapshot,
             response_snapshot, status, started_at, completed_at,
             total_questions, correct_count, total_time_ms, created_at, updated_at)
        VALUES (%s, %s, %s, %s, %s, 'completed', %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            attempt["id"],
            attempt["student_id"],
            attempt["blueprint_revision"],
            Jsonb(attempt["question_snapshot"]),
            Jsonb(attempt["response_snapshot"]),
            attempt["started_at"],
            attempt["completed_at"],
            attempt["total_questions"],
            attempt["correct_count"],
            attempt["total_time_ms"],
            attempt["started_at"],
            attempt["completed_at"],
        ),
    )


def apply_import(db: Any, plan: Mapping[str, Any]) -> dict[str, int]:
    state = _existing_state(db, plan)
    created_students = 0
    created_attempts = 0
    created_responses = 0
    rows_by_session: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in plan["response_rows"]:
        rows_by_session[row["quiz_id"]].append(row)

    for student in plan["students"]:
        if student["id"] not in state["students"]:
            _insert_student(db, student)
            created_students += 1

    for attempt in plan["attempts"]:
        if attempt["id"] in state["skip_sessions"]:
            continue
        _insert_attempt(db, attempt)
        upsert_raw_responses(db, rows_by_session[attempt["id"]])
        created_attempts += 1
        created_responses += len(rows_by_session[attempt["id"]])

    for student in plan["students"]:
        rebuild_student_vocabulary_mastery(db, student["id"])
    return {
        "created_students": created_students,
        "created_attempts": created_attempts,
        "created_responses": created_responses,
        "rebuilt_students": len(plan["students"]),
    }


def _print_summary(summary: Mapping[str, Any]) -> None:
    print(
        "Validated {student_count} students x {question_count} questions "
        "= {response_count} responses [OK]".format(**summary)
    )
    print(
        "Correct: {correct_count}; incorrect: {incorrect_count}; "
        "accuracy: {accuracy:.1%}".format(
            accuracy=summary["correct_count"] / summary["response_count"]
            if summary["response_count"]
            else 0.0,
            **summary,
        )
    )
    for mode, correct in sorted(summary["correct_by_mode"].items()):
        print(f"  {mode}: {correct} correct")
    print(f"Canonicalized story aliases: {len(summary['source_aliases'])}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="Validate only; this is the default.")
    mode.add_argument("--apply", action="store_true", help="Write the validated import to the database.")
    parser.add_argument("--database-url", default=os.environ.get("DATABASE_URL", ""))
    parser.add_argument(
        "--replace",
        action="store_true",
        help="Delete this importer's earlier synthetic data for the workbook's students, then import "
        "(use after the active blueprint changed). Same transaction as the import.",
    )
    parser.add_argument("--force", action="store_true", help="Allow a non-localhost database URL.")
    args = parser.parse_args(argv)
    if not args.workbook.is_file():
        parser.error(f"Workbook not found: {args.workbook}")
    if not args.database_url:
        parser.error("No database URL. Pass --database-url or set DATABASE_URL.")
    if not _looks_local(args.database_url) and not args.force:
        parser.error("Refusing to touch a non-localhost database without --force.")

    import db

    if args.database_url != db.DATABASE_URL:
        db.reset_pool_for_tests(args.database_url)
    imported_at = _now_iso()
    try:
        with connect_db() as connection:
            blueprint = load_active_blueprint(connection)
            rows = read_workbook(args.workbook)
            plan = build_import_plan(rows, blueprint, imported_at=imported_at)
            _print_summary(plan["summary"])
            if args.replace:
                previous = previous_import_counts(connection, plan)
                print(
                    f"--replace: {previous['attempts']} earlier imported attempts and "
                    f"{previous['responses']} responses will be deleted first."
                )
                if not args.apply:
                    print("Validation only; no database changes made. Add --apply to replace.")
                    return 0
                deleted = delete_previous_import(connection, plan)
                print(f"Deleted {deleted['attempts']} attempts and {deleted['responses']} responses.")
            state = _existing_state(connection, plan)
            print(
                f"Existing compatible sessions: {len(state['skip_sessions'])}; "
                f"new sessions: {len(plan['attempts']) - len(state['skip_sessions'])}."
            )
            if not args.apply:
                print("Validation only; no database changes made.")
                return 0
            result = apply_import(connection, plan)
        print(
            "Applied: {created_students} students, {created_attempts} placement attempts, "
            "{created_responses} responses; rebuilt {rebuilt_students} mastery projections.".format(**result)
        )
        return 0
    finally:
        db.close_db()


if __name__ == "__main__":
    raise SystemExit(main())
