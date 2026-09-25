"""Placement-test import, snapshot, grading, and BKT evidence orchestration."""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from typing import Any
from uuid import uuid4

import openpyxl

from analytics.bkt_mastery import rebuild_student_vocabulary_mastery, upsert_raw_responses
from domain.vocabulary.assessment import (
    ANSWER_FORMAT_BY_ROUND,
    QUESTION_TYPE_BY_ROUND,
    TIER_BY_ROUND,
    normalize_answer,
    numeric_to_tone_marked,
)
from repositories import placement_test_repository as repo


PLACEMENT_RESOLVER_VERSION = "placement-assessment-v1"
SUPPORTED_TYPES = {
    "basic_meaning_mcq": (1, "tier1", "meaning"),
    "character_to_pinyin_typing": (2, "tier2", "pinyin_production"),
    "context_cloze_mcq": (3, "tier3", "contextual_recall"),
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if suffix in {"xlsx", "xlsm"}:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = next(
                (workbook[name] for name in workbook.sheetnames if name.strip().casefold() in {"questions", "placement"}),
                workbook[workbook.sheetnames[0]],
            )
            iterator = sheet.iter_rows(values_only=True)
            try:
                header = [_cell_text(value) for value in next(iterator)]
            except StopIteration:
                raise ValueError("File has no header row.") from None
            if header not in (["Word Key", "Round"], ["Word Key", "Question Type"]):
                raise ValueError("File must contain exactly two columns: 'Word Key' and 'Round' or 'Question Type'.")
            return [
                {header[0]: _cell_text(row[0]), header[1]: _cell_text(row[1])}
                for row in iterator
                if len(row) >= 2 and _cell_text(row[0]) and _cell_text(row[1])
            ]
        finally:
            workbook.close()
    if suffix == "xls":
        raise ValueError("The legacy .xls format is not supported - save as .xlsx or .csv and re-upload.")
    if suffix != "csv":
        raise ValueError("Placement import accepts CSV or XLSX files only.")
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    fieldnames = [str(value).strip() for value in (reader.fieldnames or [])]
    if fieldnames not in (["Word Key", "Round"], ["Word Key", "Question Type"]):
        raise ValueError("File must contain exactly two columns: 'Word Key' and 'Round' or 'Question Type'.")
    return [
        {fieldnames[0]: str(row.get(fieldnames[0]) or "").strip(), fieldnames[1]: str(row.get(fieldnames[1]) or "").strip()}
        for row in reader
        if str(row.get(fieldnames[0]) or "").strip() and str(row.get(fieldnames[1]) or "").strip()
    ]


def _round_from_value(value: str) -> int:
    normalized = value.strip().casefold()
    if normalized in {question_type.casefold() for question_type in QUESTION_TYPE_BY_ROUND.values()}:
        return next(
            round_number
            for round_number, question_type in QUESTION_TYPE_BY_ROUND.items()
            if question_type.casefold() == normalized
        )
    normalized = normalized.replace("round", "").strip()
    if normalized in {"1", "2", "3"}:
        return int(normalized)
    raise ValueError(f"Unsupported round or question type: {value}")


def _public_question(question: dict[str, Any]) -> dict[str, Any]:
    return {
        key: question[key]
        for key in (
            "questionId", "sourceStoryId", "sourceStoryTitle", "sourceWordId", "round", "tier",
            "position", "questionType", "answerFormat", "targetWord", "prompt", "options", "audioUrl",
        )
        if key in question
    }


def _load_candidates(db: Any) -> tuple[dict[tuple[str, int], list[dict[str, Any]]], dict[tuple[str, int], list[dict[str, Any]]]]:
    published: dict[tuple[str, int], list[dict[str, Any]]] = {}
    all_stories: dict[tuple[str, int], list[dict[str, Any]]] = {}
    rows = db.execute("SELECT id, title, published, vocab_assessment FROM custom_stories").fetchall()
    for story in rows:
        assessment = story.get("vocab_assessment")
        if not isinstance(assessment, list):
            continue
        target = published if story.get("published") else all_stories
        for item in assessment:
            if not isinstance(item, dict):
                continue
            word_key = str(item.get("wordId") or "").strip()
            question_type = str(item.get("questionType") or "").strip()
            round_number = item.get("round")
            if round_number is None:
                round_number = next((number for number, kind in QUESTION_TYPE_BY_ROUND.items() if kind == question_type), None)
            try:
                round_number = int(round_number)
            except (TypeError, ValueError):
                continue
            if not word_key or round_number not in TIER_BY_ROUND:
                continue
            target.setdefault((word_key, round_number), []).append({
                "storyId": str(story["id"]),
                "storyTitle": str(story.get("title") or story["id"]),
                "item": item,
            })
    # Include published rows in the all-stories view so an unpublished-only
    # diagnostic can distinguish unknown from unpublished content.
    for key, matches in published.items():
        all_stories.setdefault(key, []).extend(matches)
    return published, all_stories


def _snapshot_question(word_key: str, round_number: int, match: dict[str, Any], position: int) -> dict[str, Any]:
    item = match["item"]
    question_type = str(item.get("questionType") or "")
    facts = SUPPORTED_TYPES.get(question_type)
    if facts is None:
        raise ValueError(f"Word Key {word_key}, Round {round_number} uses unsupported question type {question_type!r}.")
    if round_number != facts[0]:
        raise ValueError(f"Word Key {word_key}, Round {round_number} uses {question_type}, expected {QUESTION_TYPE_BY_ROUND[round_number]}.")
    answer_format = str(item.get("answerFormat") or "")
    expected_format = "free_text" if question_type == "character_to_pinyin_typing" else "single_choice"
    if answer_format != expected_format:
        raise ValueError(f"Word Key {word_key}, Round {round_number} has unsupported answer format {answer_format!r}.")
    accepted = item.get("acceptedAnswers") or []
    question_id = str(item.get("questionId") or "").strip()
    if not question_id:
        raise ValueError(f"Word Key {word_key}, Round {round_number} has no source questionId in the published bank.")
    return {
        "questionId": question_id,
        "sourceStoryId": match["storyId"],
        "sourceStoryTitle": match["storyTitle"],
        "sourceWordId": str(item.get("wordId") or ""),
        "round": round_number,
        "tier": TIER_BY_ROUND[round_number],
        "position": position,
        "questionType": question_type,
        "answerFormat": answer_format,
        "targetWord": str(item.get("targetWord") or ""),
        "pinyin": str(item.get("pinyin") or ""),
        "prompt": str(item.get("prompt") or ""),
        "options": [str(value) for value in (item.get("options") or [])],
        "correctAnswer": str(item.get("correctAnswer") or ""),
        "acceptedAnswers": [str(value) for value in accepted],
        "explanation": str(item.get("explanation") or ""),
        "audioUrl": item.get("audioUrl") if isinstance(item.get("audioUrl"), str) else None,
    }


def _validate_rows(db: Any, rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    if not rows:
        raise ValueError("File has no data rows.")
    normalized: list[tuple[str, int]] = []
    for row in rows:
        word_key = row.get("Word Key", "").strip()
        selector = row.get("Round") or row.get("Question Type") or ""
        try:
            normalized.append((word_key, _round_from_value(selector)))
        except ValueError as exc:
            raise ValueError(f"{word_key or '<empty>'}: {exc}") from exc
    seen: set[tuple[str, int]] = set()
    duplicates: list[tuple[str, int]] = []
    for key in normalized:
        if key in seen and key not in duplicates:
            duplicates.append(key)
        seen.add(key)
    if duplicates:
        raise ValueError("Duplicate Word Key + Round: " + ", ".join(f"{word} + {round_number}" for word, round_number in duplicates[:20]))
    published, all_stories = _load_candidates(db)
    snapshot: list[dict[str, Any]] = []
    for position, (word_key, round_number) in enumerate(normalized, start=1):
        candidates = published.get((word_key, round_number), [])
        if not candidates:
            if all_stories.get((word_key, round_number)):
                raise ValueError(f"Word Key {word_key}, Round {round_number} belongs to an unpublished story.")
            raise ValueError(f"Unknown Word Key + Round: {word_key} + {round_number}")
        if len(candidates) > 1:
            stories = ", ".join(sorted({str(candidate["storyId"]) for candidate in candidates}))
            raise ValueError(f"Ambiguous Word Key + Round {word_key} + {round_number}; found in published stories: {stories}.")
        snapshot.append(_snapshot_question(word_key, round_number, candidates[0], position))
    return snapshot


def build_preview(db: Any, content: bytes, filename: str) -> dict[str, Any]:
    try:
        rows = _parse_rows(filename, content)
        questions = _validate_rows(db, rows)
        errors: list[str] = []
    except ValueError as exc:
        questions = []
        errors = [str(exc)]
    current = repo.get_active_blueprint(db)
    return {
        "valid": not errors,
        "rowIssues": errors,
        "questionCount": len(questions),
        "questions": [_public_question(question) for question in questions],
        "currentRevision": current["revision"] if current else None,
    }


def replace_from_upload(db: Any, content: bytes, filename: str) -> dict[str, Any]:
    rows = _parse_rows(filename, content)
    questions = _validate_rows(db, rows)
    blueprint = repo.replace_active_blueprint(db, questions, _now())
    return {
        "configured": True,
        "revision": blueprint["revision"],
        "questionCount": len(questions),
        "questions": [_public_question(question) for question in questions],
    }


def get_student_blueprint(db: Any) -> dict[str, Any]:
    blueprint = repo.get_active_blueprint(db)
    if not blueprint or not blueprint.get("questions"):
        return {"configured": False, "revision": None, "questionCount": 0, "questions": []}
    questions = blueprint["questions"]
    return {
        "configured": True,
        "revision": blueprint["revision"],
        "questionCount": len(questions),
        "questions": [_public_question(question) for question in questions],
    }


def get_admin_blueprint(db: Any) -> dict[str, Any]:
    blueprint = repo.get_active_blueprint(db)
    if not blueprint:
        return {"configured": False, "revision": None, "questionCount": 0, "questions": []}
    return {
        "configured": bool(blueprint.get("questions")),
        "revision": blueprint["revision"],
        "questionCount": len(blueprint.get("questions") or []),
        "questions": blueprint.get("questions") or [],
        "updatedAt": blueprint.get("updated_at"),
    }


def start_attempt(db: Any, student_id: str) -> dict[str, Any]:
    blueprint = repo.get_active_blueprint(db)
    questions = blueprint.get("questions") if blueprint else None
    if not questions:
        raise ValueError("Placement test is not configured.")
    started_at = _now()
    attempt_id = f"placement-{student_id}-{uuid4().hex[:12]}"
    repo.insert_attempt(
        db,
        attempt_id=attempt_id,
        student_id=student_id,
        blueprint_revision=int(blueprint["revision"]),
        question_snapshot=questions,
        started_at=started_at,
    )
    return {
        "attemptId": attempt_id,
        "revision": blueprint["revision"],
        "totalQuestions": len(questions),
        "questions": [_public_question(question) for question in questions],
    }


def _parse_answered_at(value: Any, fallback: str) -> str:
    if not isinstance(value, str) or not value.strip():
        return fallback
    try:
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError:
        return fallback
    if parsed.tzinfo is None:
        return fallback
    return parsed.astimezone(timezone.utc).isoformat()


def _grade(question: dict[str, Any], selected_answer: str) -> bool:
    candidates = [selected_answer]
    if question["questionType"] == "character_to_pinyin_typing":
        candidates.append(numeric_to_tone_marked(selected_answer))
    accepted = question.get("acceptedAnswers") or [question.get("correctAnswer", "")]
    return any(
        normalize_answer(candidate) == normalize_answer(answer)
        for candidate in candidates
        for answer in accepted
    )


def complete_attempt(db: Any, student_id: str, attempt_id: str, responses: list[dict[str, Any]], completed_at: str | None) -> dict[str, Any]:
    attempt = repo.get_attempt_for_update(db, attempt_id, student_id)
    if not attempt:
        raise LookupError("Placement attempt was not found.")
    if attempt["status"] == "completed":
        correct_count = int(attempt.get("correct_count") or 0)
        total = int(attempt["total_questions"])
        return _result(attempt_id, total, correct_count)
    questions = attempt["question_snapshot"] or []
    by_id = {question["questionId"]: question for question in questions}
    if len(responses) != len(questions):
        raise ValueError(f"Expected exactly {len(questions)} responses.")
    seen: set[str] = set()
    normalized_responses: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []
    total_time_ms = 0
    fallback_time = _parse_answered_at(completed_at, _now()) if completed_at else _now()
    for raw in responses:
        question_id = str(raw.get("questionId") or "").strip()
        selected_answer = raw.get("selectedAnswer")
        if question_id in seen or question_id not in by_id:
            raise ValueError(f"Response has unknown or duplicate Question ID: {question_id or '<empty>'}.")
        if not isinstance(selected_answer, str) or not selected_answer.strip():
            raise ValueError(f"Question {question_id} needs a selected answer.")
        time_ms = raw.get("timeMs", 0)
        if not isinstance(time_ms, int) or time_ms < 0:
            raise ValueError(f"Question {question_id} has invalid timeMs.")
        question = by_id[question_id]
        answered_at = _parse_answered_at(raw.get("answeredAt"), fallback_time)
        correct = _grade(question, selected_answer)
        total_time_ms += time_ms
        seen.add(question_id)
        normalized_responses.append({
            "questionId": question_id,
            "selectedAnswer": selected_answer,
            "timeMs": time_ms,
            "answeredAt": answered_at,
            "correct": correct,
            "position": question["position"],
        })
    if seen != set(by_id):
        missing = sorted(set(by_id) - seen)
        raise ValueError("Missing responses for: " + ", ".join(missing[:20]))
    normalized_responses.sort(key=lambda response: response["position"])
    correct_count = sum(1 for response in normalized_responses if response["correct"])
    for response in normalized_responses:
        question = by_id[response["questionId"]]
        round_number, tier, dimension = SUPPORTED_TYPES[question["questionType"]]
        question_prompt = question["prompt"]
        if question["questionType"] == "character_to_pinyin_typing" and not question_prompt:
            question_prompt = f"Type the pinyin for {question['targetWord']}."
        rows.append({
            "student_id": student_id,
            "word_id": question["sourceWordId"],
            "word": question["targetWord"],
            "lesson_id": question["sourceStoryId"],
            "quiz_id": attempt_id,
            "attempt_id": attempt_id,
            "item_id": question["questionId"],
            "question_type": question["questionType"],
            "round_type": str(round_number),
            "knowledge_dimension": dimension,
            "activity_type": "diagnostic",
            "diagnostic_exposure_id": f"placement:{attempt_id}:{question['questionId']}",
            "bkt_eligible": True,
            "bkt_eligibility_errors": [],
            "selected_answer": response["selectedAnswer"],
            "correct_answer": question["correctAnswer"],
            "presented_options": question["options"],
            "question_prompt": question_prompt,
            "answered_at": response["answeredAt"],
            "correct": response["correct"],
            "response_time_ms": response["timeMs"],
            "occurred_at": response["answeredAt"],
            "occurred_at_utc": datetime.fromisoformat(response["answeredAt"].replace("Z", "+00:00")).astimezone(timezone.utc),
            "evidence_origin": "real",
            "resolver_version": PLACEMENT_RESOLVER_VERSION,
            "attempt_order": question["position"] - 1,
            "quiz_level": tier,
            "quiz_mode": tier,
            "research_study_id": None,
        })
    upsert_raw_responses(db, rows)
    rebuild_student_vocabulary_mastery(db, student_id)
    finished_at = _parse_answered_at(completed_at, _now())
    repo.complete_attempt(
        db,
        attempt_id=attempt_id,
        response_snapshot=normalized_responses,
        completed_at=finished_at,
        correct_count=correct_count,
        total_time_ms=total_time_ms,
    )
    return _result(attempt_id, len(questions), correct_count)


def _result(attempt_id: str, total: int, correct: int) -> dict[str, Any]:
    percentage = round((correct / total) * 100) if total else 0
    return {
        "attemptId": attempt_id,
        "totalQuestions": total,
        "correctCount": correct,
        "percentage": percentage,
        "messageKey": "CONGRATULATIONS" if percentage >= 70 else "KEEP_PRACTICING",
    }
