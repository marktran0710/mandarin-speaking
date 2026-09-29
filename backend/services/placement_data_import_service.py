"""Admin import orchestration for placement response workbooks."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

import security.auth as auth
from scripts import import_placement_bkt_workbook as workbook_import


SUPPORTED_SUFFIXES = {"xlsx", "xlsm"}
SAMPLE_TIMESTAMP = "2026-09-25T00:00:00+00:00"


def _validate_filename(filename: str) -> None:
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Placement response import accepts XLSX or XLSM files only.")


def _apply_result(result: dict[str, int]) -> dict[str, int]:
    # apply_import() is shared with the CLI script's own snake_case
    # formatting, so translate here rather than changing its return shape.
    return {
        "createdStudents": result["created_students"],
        "createdAttempts": result["created_attempts"],
        "createdResponses": result["created_responses"],
        "rebuiltStudents": result["rebuilt_students"],
    }


def _summary(plan: dict[str, Any], state: dict[str, Any]) -> dict[str, Any]:
    source_aliases = plan["summary"].get("source_aliases") or {}
    return {
        "studentCount": int(plan["summary"]["student_count"]),
        "questionCount": int(plan["summary"]["question_count"]),
        "responseCount": int(plan["summary"]["response_count"]),
        "correctCount": int(plan["summary"]["correct_count"]),
        "incorrectCount": int(plan["summary"]["incorrect_count"]),
        "correctByMode": dict(plan["summary"].get("correct_by_mode") or {}),
        "canonicalizedStoryAliases": len(source_aliases),
        "existingSessions": len(state["skip_sessions"]),
        "newSessions": len(plan["attempts"]) - len(state["skip_sessions"]),
    }


def _build_plan(db: Any, content: bytes, filename: str) -> tuple[dict[str, Any], dict[str, Any]]:
    _validate_filename(filename)
    if not content:
        raise ValueError("The placement response workbook is empty.")
    blueprint = workbook_import.load_active_blueprint(db)
    try:
        rows = workbook_import.read_workbook(BytesIO(content))
    except (InvalidFileException, OSError, ValueError) as exc:
        raise ValueError(f"Could not read the placement response workbook: {exc}") from exc
    plan = workbook_import.build_import_plan(
        rows,
        blueprint,
        imported_at=datetime.now(timezone.utc).isoformat(),
        expected_student_ids=workbook_import.EXPECTED_STUDENT_IDS,
    )
    return plan, workbook_import._existing_state(db, plan)


def preview_import(db: Any, content: bytes, filename: str) -> dict[str, Any]:
    try:
        plan, state = _build_plan(db, content, filename)
    except ValueError as exc:
        return {
            "valid": False,
            "filename": filename,
            "rowIssues": [str(exc)],
            "studentCount": 0,
            "questionCount": 0,
            "responseCount": 0,
            "correctCount": 0,
            "incorrectCount": 0,
            "correctByMode": {},
            "canonicalizedStoryAliases": 0,
            "existingSessions": 0,
            "newSessions": 0,
        }
    return {
        "valid": True,
        "filename": filename,
        "rowIssues": [],
        **_summary(plan, state),
    }


def confirm_import(db: Any, content: bytes, filename: str) -> dict[str, Any]:
    plan, state = _build_plan(db, content, filename)
    result = workbook_import.apply_import(db, plan)
    return {
        "valid": True,
        "filename": filename,
        **_summary(plan, state),
        **_apply_result(result),
    }


def replace_import(db: Any, content: bytes, filename: str) -> dict[str, Any]:
    """Delete this importer's earlier data for the workbook's students, then import fresh.

    Mirrors ``import_placement_bkt_workbook.py --replace``: only rows this
    importer itself wrote (``evidence_origin='synthetic'`` and its own
    resolver version) are ever deleted, and only attempts holding exclusively
    those rows - real evidence and other attempts are never touched, and
    student accounts are kept so their ids stay stable. This is what clears
    a stale "question snapshot differs from the active blueprint" block after
    the active blueprint changes.
    """
    _validate_filename(filename)
    if not content:
        raise ValueError("The placement response workbook is empty.")
    blueprint = workbook_import.load_active_blueprint(db)
    try:
        rows = workbook_import.read_workbook(BytesIO(content))
    except (InvalidFileException, OSError, ValueError) as exc:
        raise ValueError(f"Could not read the placement response workbook: {exc}") from exc
    plan = workbook_import.build_import_plan(
        rows,
        blueprint,
        imported_at=datetime.now(timezone.utc).isoformat(),
        expected_student_ids=workbook_import.EXPECTED_STUDENT_IDS,
    )

    deleted = workbook_import.delete_previous_import(db, plan)
    state = workbook_import._existing_state(db, plan)
    result = workbook_import.apply_import(db, plan)
    return {
        "valid": True,
        "filename": filename,
        "deletedStudents": deleted["students"],
        "deletedAttempts": deleted["attempts"],
        "deletedResponses": deleted["responses"],
        **_summary(plan, state),
        **_apply_result(result),
    }


def activate_imported_student_accounts(db: Any, temporary_password: str) -> dict[str, Any]:
    """Give the imported SIM cohort a usable shared temporary password.

    Placement imports already create student rows so every synthetic response
    has a real owner, but those rows deliberately start with an unknown random
    password and ``password_reset_required`` enabled. Activating them in place
    preserves all placement/BKT ownership while keeping the accounts marked as
    test data.
    """
    auth.validate_password_policy(temporary_password)
    expected_ids = list(workbook_import.EXPECTED_STUDENT_IDS)
    rows = db.execute(
        """
        SELECT s.id, s.name
        FROM students s
        WHERE s.id = ANY(%s)
          AND s.is_test_account = TRUE
          AND EXISTS (
              SELECT 1
              FROM vocab_quiz_responses response
              WHERE response.student_id = s.id
                AND response.evidence_origin = 'synthetic'
                AND response.resolver_version = %s
          )
        ORDER BY s.id
        FOR UPDATE
        """,
        (expected_ids, workbook_import.IMPORT_RESOLVER_VERSION),
    ).fetchall()
    found_ids = {str(row["id"]) for row in rows}
    missing_ids = [student_id for student_id in expected_ids if student_id not in found_ids]
    if missing_ids:
        raise ValueError(
            "Cannot activate placement student logins until the complete imported cohort exists. "
            f"Missing {len(missing_ids)} account(s): {', '.join(missing_ids[:8])}"
            + ("..." if len(missing_ids) > 8 else "")
        )

    password_hash = auth.hash_password(temporary_password)
    updated = db.execute(
        """
        UPDATE students
        SET password = %s,
            password_reset_required = FALSE,
            password_version = COALESCE(password_version, 0) + 1,
            status = 'active'
        WHERE id = ANY(%s)
          AND is_test_account = TRUE
        """,
        (password_hash, expected_ids),
    ).rowcount
    if updated != len(expected_ids):
        raise ValueError("The imported placement cohort changed while login accounts were being activated.")
    return {
        "activatedAccounts": updated,
        "studentIds": expected_ids,
        "testAccounts": True,
    }


def build_sample_workbook(db: Any) -> bytes:
    blueprint = workbook_import.load_active_blueprint(db)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Responses"
    sheet.append(list(workbook_import.EXPECTED_HEADERS))

    questions = blueprint["questions"]
    for student_id in workbook_import.EXPECTED_STUDENT_IDS:
        student_name = f"Synthetic Student {student_id.removeprefix('SIM')}"
        session_id = f"PLACEMENT-{student_id}-V1"
        for question in questions:
            accepted = question.get("acceptedAnswers") or [question.get("correctAnswer") or ""]
            sheet.append([
                student_id,
                student_name,
                session_id,
                str(question.get("sourceStoryId") or ""),
                str(question.get("questionId") or ""),
                str(question.get("tier") or ""),
                str(accepted[0] if accepted else ""),
                SAMPLE_TIMESTAMP,
                1000,
            ])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [16, 26, 28, 28, 16, 12, 28, 30, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Placement response workbook"])
    instructions.append(["Replace selected_answer, answered_at, and time_ms with the real response values."])
    instructions.append(["Keep the exact Responses headers and the 40-student x 28-question row coverage."])
    instructions.append(["item_id must be the questionId from the active placement blueprint."])
    instructions.append(["mode must match the question tier: tier1, tier2, or tier3."])
    instructions.column_dimensions["A"].width = 110

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
