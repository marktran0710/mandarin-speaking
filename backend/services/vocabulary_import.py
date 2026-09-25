"""Admin-facing canonical vocabulary/question import.

Reuses the exact row shape and per-row/per-word validation
``scripts/import_question_bank_workbook.py`` already established for the
canonical Chapters 5-8 workbook, minus that script's one-time assumptions
(a fixed 648-row count, chapters 5-8 only, a single hardcoded file path).
``build_payloads`` and ``find_story_for_part`` are reused as-is - neither
one carried those assumptions to begin with.

The admin import replaces each matched lesson's complete ``vocab_assessment``
bank. Existing audio is carried forward by ``wordId``. Once a lesson has a
canonical quiz bank, legacy frame/story vocabulary is cleared so quiz content
is the one vocabulary source used by the story.
Nothing is written until ``apply_vocabulary_import`` is called explicitly -
preview is read-only.
"""

from __future__ import annotations

import csv
import io
import json
from collections import defaultdict
from typing import Any

import openpyxl
from psycopg.types.json import Jsonb

from domain.vocabulary.assessment import (
    ANSWER_FORMAT_BY_ROUND,
    QUESTION_TYPE_BY_ROUND,
    ROUNDS as ASSESSMENT_ROUNDS,
    validate_assessment_payload,
)
from scripts.import_question_bank_workbook import REQUIRED_COLUMNS, ROUNDS, build_payloads
from scripts.seed_quiz_assessments import find_story_for_part
import services.media as media_service


ROUND_ALIASES = {"1": "Round 1", "2": "Round 2", "3": "Round 3"}
INPUT_MODE_ALIASES = {"mcq": "click"}
IMPORT_MODE = "replace_lesson"
RETIRED_COLUMNS = frozenset({"Source Type"})
IMPORT_REQUIRED_COLUMNS = REQUIRED_COLUMNS - RETIRED_COLUMNS
TEMPLATE_COLUMNS = (
    "Question ID", "Word Key", "Chapter", "Section", "Item",
    "Traditional Chinese", "Pinyin", "POS", "English Meaning", "Round",
    "Question Type", "Input Mode", "Prompt", "Option A", "Option B",
    "Option C", "Option D", "Correct Option", "Correct Answer", "Accepted Answers",
)


def _ensure_import_mode(mode: str) -> None:
    if mode != IMPORT_MODE:
        raise ValueError(f"Unsupported vocabulary import mode {mode!r}; expected {IMPORT_MODE!r}.")


def _normalize_import_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """Accept the compact workbook vocabulary used by the admin template.

    The canonical payload still uses the internal ``Round 1``/``click`` values,
    but exported workbooks commonly use numeric rounds and ``mcq``. Normalize
    those aliases once at the upload boundary so validation and publishing see
    one consistent shape.
    """
    normalized: list[dict[str, str]] = []
    for raw_row in rows:
        row = dict(raw_row)
        round_value = (row.get("Round") or "").strip()
        input_mode = (row.get("Input Mode") or "").strip()
        row["Round"] = ROUND_ALIASES.get(round_value.casefold(), round_value)
        row["Input Mode"] = INPUT_MODE_ALIASES.get(input_mode.casefold(), input_mode)
        normalized.append(row)
    return normalized


def parse_csv_rows(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    missing = IMPORT_REQUIRED_COLUMNS.difference(reader.fieldnames or ())
    if missing:
        raise ValueError(f"File is missing required columns: {', '.join(sorted(missing))}")
    return _normalize_import_rows([dict(row) for row in reader if row.get("Question ID")])


def _cell_text(value: object) -> str:
    """Excel stores numbers as float/int; normalize integral values before validation."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_xlsx_rows(content: bytes) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = next(
            (workbook[name] for name in workbook.sheetnames if name.strip().casefold() in {"questions", "quiz questions"}),
            workbook[workbook.sheetnames[0]],
        )
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise ValueError("File has no header row.") from None
        header = [_cell_text(cell) for cell in header_row]
        missing = IMPORT_REQUIRED_COLUMNS.difference(header)
        if missing:
            raise ValueError(f"File is missing required columns: {', '.join(sorted(missing))}")
        rows: list[dict[str, str]] = []
        for raw_row in rows_iter:
            row = {header[index]: _cell_text(raw_row[index]) for index in range(len(header)) if index < len(raw_row)}
            if row.get("Question ID"):
                rows.append(row)
        return _normalize_import_rows(rows)
    finally:
        workbook.close()


def parse_uploaded_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    """Dispatch by extension - both formats produce the identical row shape
    that validate_import_rows/build_payloads consume, so nothing downstream
    needs to know which one was uploaded."""
    suffix = (filename or "").rsplit(".", 1)[-1].casefold() if "." in (filename or "") else ""
    if suffix in {"xlsx", "xlsm"}:
        return parse_xlsx_rows(content)
    if suffix == "xls":
        raise ValueError("The legacy .xls format is not supported - save as .xlsx or .csv and re-upload.")
    return parse_csv_rows(content)


def _answers(value: str) -> list[str]:
    return [part.strip() for part in (value or "").split("|") if part.strip()]


def validate_import_rows(rows: list[dict[str, str]]) -> list[str]:
    """Same per-row/per-word checks as the workbook importer, generalized to
    any section/row count (no fixed 648-row or chapter-5-8 assumption)."""
    issues: list[str] = []
    if not rows:
        issues.append("File has no data rows.")
        return issues
    seen_question_ids: set[str] = set()
    by_word: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        qid = row.get("Question ID", "")
        if qid in seen_question_ids:
            issues.append(f"duplicate Question ID: {qid}")
        seen_question_ids.add(qid)
        by_word[row.get("Word Key", "")].append(row)
        section = row.get("Section", "")
        if "-" not in section or not all(part.strip().isdigit() for part in section.split("-", 1)):
            issues.append(f"{qid}: Section must look like '5-1' (chapter-part), found {section!r}")
        round_number = ROUNDS.get(row.get("Round", ""))
        if round_number is None:
            issues.append(f"{qid}: unsupported round {row.get('Round')!r}")
            continue
        answer_format = ANSWER_FORMAT_BY_ROUND[round_number]
        expected_type = QUESTION_TYPE_BY_ROUND[round_number]
        if row.get("Question Type") != expected_type:
            issues.append(f"{qid}: {row.get('Round')} uses {row.get('Question Type')}, expected {expected_type}")
        if row.get("Input Mode") == "click" and answer_format != "single_choice":
            issues.append(f"{qid}: click input is not a single-choice question")
        if row.get("Input Mode") == "free_text" and answer_format != "free_text":
            issues.append(f"{qid}: free-text input has the wrong answer format")
        if answer_format == "single_choice":
            options = [row.get(f"Option {letter}", "") for letter in "ABCD"]
            if len(set(options)) != 4 or any(not option for option in options):
                issues.append(f"{qid}: options must contain four distinct values")
            correct_option = row.get("Correct Option", "")
            if correct_option not in "ABCD":
                issues.append(f"{qid}: invalid correct option {correct_option!r}")
            elif options["ABCD".index(correct_option)] != row.get("Correct Answer"):
                issues.append(f"{qid}: correct option does not match correct answer")
        else:
            accepted = _answers(row.get("Accepted Answers", ""))
            if row.get("Correct Answer") not in accepted:
                issues.append(f"{qid}: pinyin correct answer is not accepted")
        for required in ("Word Key", "Traditional Chinese", "Pinyin", "POS", "English Meaning", "Prompt", "Correct Answer"):
            if not (row.get(required) or "").strip():
                issues.append(f"{qid}: empty {required}")
    for word_id, word_rows in by_word.items():
        if len(word_rows) != len(ASSESSMENT_ROUNDS) or {
            ROUNDS.get(row.get("Round", "")) for row in word_rows
        } != set(ASSESSMENT_ROUNDS):
            issues.append(f"{word_id}: expected exactly one row for each round (Round 1/2/3)")
        for field in ("Chapter", "Section", "Traditional Chinese", "Pinyin", "POS", "English Meaning"):
            if len({row.get(field) for row in word_rows}) != 1:
                issues.append(f"{word_id}: inconsistent {field} metadata across its three rows")
    return issues


def _canonical_payloads(rows: list[dict[str, str]]) -> dict[str, list[dict[str, Any]]]:
    """Build assessment rows and remove retired metadata from new payloads."""
    payloads = build_payloads(rows)
    return {
        section: [
            {key: value for key, value in question.items() if key not in {"sourceType"}}
            for question in payload
        ]
        for section, payload in payloads.items()
    }


def _audio_by_word(assessment: list[dict[str, Any]]) -> dict[str, str]:
    audio: dict[str, str] = {}
    for question in assessment:
        word_id = str(question.get("wordId") or "").strip()
        url = question.get("audioUrl")
        if word_id and word_id not in audio and isinstance(url, str) and url.strip():
            audio[word_id] = url.strip()
    return audio


def _replace_assessment(
    existing: list[dict[str, Any]], incoming: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], set[str], set[str], set[str]]:
    """Replace a lesson bank while preserving per-word audio references."""
    existing_word_ids = {str(question.get("wordId") or "").strip() for question in existing if question.get("wordId")}
    incoming_word_ids = {str(question.get("wordId") or "").strip() for question in incoming if question.get("wordId")}
    audio_by_word = _audio_by_word(existing)
    preserved_audio = incoming_word_ids & audio_by_word.keys()
    missing_audio = incoming_word_ids - preserved_audio
    replacement: list[dict[str, Any]] = []
    for question in incoming:
        row = dict(question)
        audio_url = audio_by_word.get(str(row.get("wordId") or "").strip())
        if audio_url:
            row["audioUrl"] = audio_url
        replacement.append(row)
    return replacement, incoming_word_ids - existing_word_ids, existing_word_ids - incoming_word_ids, missing_audio


LEGACY_VOCABULARY_FRAME_FIELDS = (
    "vocabulary",
    "vocabularyGroups",
    "vocabularyPinyin",
    "vocabularyPos",
    "vocabularyTranslation",
    "vocabularyAudioUrls",
    "vocabularyReferenceCurves",
    "vocabularyMedium",
    "vocabularyHard",
    "vocabularyPinyinMedium",
    "vocabularyPinyinHard",
    "vocabularyPosMedium",
    "vocabularyPosHard",
    "vocabularyTranslationMedium",
    "vocabularyTranslationHard",
    "vocabularyAudioUrlsMedium",
    "vocabularyAudioUrlsHard",
    "vocabularyReferenceCurvesMedium",
    "vocabularyReferenceCurvesHard",
)

LEGACY_VOCABULARY_AUDIO_FIELDS = {
    "vocabularyAudioUrls",
    "vocabularyAudioUrlsMedium",
    "vocabularyAudioUrlsHard",
}


def _frame_audio_urls(frames: object) -> set[str]:
    """Return stored per-word frame audio URLs before legacy vocab cleanup."""
    urls: set[str] = set()
    if not isinstance(frames, list):
        return urls
    for frame in frames:
        if not isinstance(frame, dict):
            continue
        for field in LEGACY_VOCABULARY_AUDIO_FIELDS:
            raw_urls = frame.get(field)
            if isinstance(raw_urls, str):
                try:
                    raw_urls = json.loads(raw_urls)
                except json.JSONDecodeError:
                    raw_urls = []
            if isinstance(raw_urls, list):
                urls.update(
                    value.strip()
                    for value in raw_urls
                    if isinstance(value, str) and value.strip().startswith("/uploads/")
                )
            elif isinstance(raw_urls, dict):
                for value in raw_urls.values():
                    if isinstance(value, str) and value.strip().startswith("/uploads/"):
                        urls.add(value.strip())
                    elif isinstance(value, list):
                        urls.update(
                            item.strip()
                            for item in value
                            if isinstance(item, str) and item.strip().startswith("/uploads/")
                        )
    return urls


def _clear_legacy_vocabulary(frames: object) -> tuple[list[dict[str, Any]], set[str]]:
    """Clear old authored vocabulary while preserving scene prompts/media."""
    normalized_frames: list[dict[str, Any]] = []
    old_audio_urls = _frame_audio_urls(frames)
    for raw_frame in frames if isinstance(frames, list) else []:
        frame = dict(raw_frame) if isinstance(raw_frame, dict) else {}
        for field in LEGACY_VOCABULARY_FRAME_FIELDS:
            if field not in frame and field not in LEGACY_VOCABULARY_FRAME_FIELDS[:7]:
                continue
            if field == "vocabulary":
                # The API's frame contract requires this field to be a string.
                frame[field] = ""
            elif field == "vocabularyGroups":
                frame[field] = []
            else:
                frame[field] = ""
        normalized_frames.append(frame)
    return normalized_frames, old_audio_urls


def _clear_legacy_vocabulary_for_lessons(db: Any, lesson_numbers: set[int]) -> set[str]:
    """Clear scene/story vocabulary on every story in imported lessons.

    This includes published teacher-created duplicates with no lesson part;
    their publication state is intentionally left unchanged.
    """
    if not lesson_numbers:
        return set()
    rows = db.execute(
        "SELECT id, frames FROM custom_stories WHERE lesson_number = ANY(%s) FOR UPDATE",
        (list(lesson_numbers),),
    ).fetchall()
    cleanup_urls: set[str] = set()
    for row in rows:
        frames, old_audio_urls = _clear_legacy_vocabulary(row.get("frames"))
        cleanup_urls.update(old_audio_urls)
        db.execute(
            """
            UPDATE custom_stories
            SET frames = %s::jsonb, story_vocabulary = NULL
            WHERE id = %s
            """,
            (Jsonb(frames), row["id"]),
        )
    return cleanup_urls


def _assessment_audio_urls(assessment: list[dict[str, Any]]) -> set[str]:
    return {
        question["audioUrl"]
        for question in assessment
        if isinstance(question, dict)
        and isinstance(question.get("audioUrl"), str)
        and question["audioUrl"].startswith("/uploads/")
    }


def _remove_unreferenced_audio(db: Any, candidates: set[str]) -> None:
    if not candidates:
        return
    rows = db.execute("SELECT vocab_assessment, frames FROM custom_stories").fetchall()
    referenced: set[str] = set()
    for row in rows:
        assessment = row.get("vocab_assessment")
        if isinstance(assessment, list):
            referenced.update(_assessment_audio_urls(assessment))
        referenced.update(_frame_audio_urls(row.get("frames")))
    for url in candidates - referenced:
        media_service.remove_uploaded_file(url)


def _section_preview(db: Any, section: str, payload: list[dict[str, Any]]) -> dict[str, Any]:
    try:
        story = find_story_for_part(db, section)
    except LookupError as exc:
        return {
            "section": section, "storyId": None, "storyTitle": None, "found": False,
            "error": str(exc), "newWords": 0, "updatedWords": 0, "removedWords": 0,
            "preservedAudio": 0, "missingAudio": 0, "questionCount": len(payload), "issues": [],
        }
    row = db.execute("SELECT vocab_assessment FROM custom_stories WHERE id = %s", (story["id"],)).fetchone()
    existing = row["vocab_assessment"] if row and isinstance(row.get("vocab_assessment"), list) else []
    existing_word_ids = {question.get("wordId") for question in existing}
    incoming_word_ids = {question["wordId"] for question in payload}
    replacement, new_words, removed_words, missing_audio = _replace_assessment(existing, payload)
    issues = validate_assessment_payload(replacement)
    return {
        "section": section,
        "storyId": story["id"],
        "storyTitle": story["title"],
        "found": True,
        "newWords": len(new_words),
        "updatedWords": len(incoming_word_ids & existing_word_ids),
        "removedWords": len(removed_words),
        "preservedAudio": len(incoming_word_ids - missing_audio),
        "missingAudio": len(missing_audio),
        "questionCount": len(payload),
        "issues": [f"{issue.code}: {issue.message}" for issue in issues],
    }


def preview_vocabulary_import(
    db: Any, content: bytes, filename: str = "", *, mode: str = IMPORT_MODE,
) -> dict[str, Any]:
    """Read-only: parse, validate, and report what an import would do."""
    _ensure_import_mode(mode)
    rows = parse_uploaded_rows(filename, content)
    row_issues = validate_import_rows(rows)
    if row_issues:
        return {
            "mode": mode, "rows": len(rows), "rowIssues": row_issues, "sections": [],
            "newWords": 0, "updatedWords": 0, "removedWords": 0,
            "preservedAudio": 0, "missingAudio": 0,
        }
    payloads = _canonical_payloads(rows)
    with_db = [_section_preview(db, section, payload) for section, payload in sorted(payloads.items())]
    return {
        "mode": mode,
        "rows": len(rows),
        "rowIssues": [],
        "sections": with_db,
        "newWords": sum(section["newWords"] for section in with_db),
        "updatedWords": sum(section["updatedWords"] for section in with_db),
        "removedWords": sum(section["removedWords"] for section in with_db),
        "preservedAudio": sum(section["preservedAudio"] for section in with_db),
        "missingAudio": sum(section["missingAudio"] for section in with_db),
    }


def apply_vocabulary_import(
    db: Any, content: bytes, filename: str = "", *, mode: str = IMPORT_MODE,
) -> dict[str, Any]:
    """Write path. Re-validates from scratch - never trusts a client-held
    preview result as proof the file is still valid to write."""
    _ensure_import_mode(mode)
    rows = parse_uploaded_rows(filename, content)
    row_issues = validate_import_rows(rows)
    if row_issues:
        raise ValueError("Row validation failed:\n" + "\n".join(row_issues[:30]))
    payloads = _canonical_payloads(rows)
    published: list[dict[str, Any]] = []
    cleanup_urls: set[str] = set()
    totals = {"newWords": 0, "updatedWords": 0, "removedWords": 0, "preservedAudio": 0, "missingAudio": 0}
    for section, payload in sorted(payloads.items()):
        story = find_story_for_part(db, section)
        row = db.execute(
            "SELECT vocab_assessment FROM custom_stories WHERE id = %s FOR UPDATE", (story["id"],),
        ).fetchone()
        existing = row["vocab_assessment"] if row and isinstance(row.get("vocab_assessment"), list) else []
        replacement, new_words, removed_words, missing_audio = _replace_assessment(existing, payload)
        issues = validate_assessment_payload(replacement)
        if issues:
            rendered = "; ".join(f"{issue.code}: {issue.message}" for issue in issues[:8])
            raise ValueError(f"{section} would fail validation after replace ({len(issues)} issues): {rendered}")
        cleanup_urls.update(
            url for question in existing
            if question.get("wordId") in removed_words
            for url in [question.get("audioUrl")]
            if isinstance(url, str) and url.startswith("/uploads/")
        )
        db.execute(
            "UPDATE custom_stories SET vocab_assessment = %s::jsonb WHERE id = %s",
            (Jsonb(replacement), story["id"]),
        )
        section_updated = len({question["wordId"] for question in payload} & {question.get("wordId") for question in existing})
        section_preserved_audio = len({question["wordId"] for question in replacement if question.get("audioUrl")})
        section_missing_audio = len({question["wordId"] for question in replacement if not question.get("audioUrl")})
        totals["newWords"] += len(new_words)
        totals["updatedWords"] += section_updated
        totals["removedWords"] += len(removed_words)
        totals["preservedAudio"] += section_preserved_audio
        totals["missingAudio"] += section_missing_audio
        published.append({
            "section": section, "storyId": story["id"], "storyTitle": story["title"],
            "questionCount": len(payload), "newWords": len(new_words), "updatedWords": section_updated,
            "removedWords": len(removed_words), "preservedAudio": section_preserved_audio,
            "missingAudio": section_missing_audio,
        })
    # The imported quiz bank is now the sole vocabulary source for every
    # story in the affected lesson(s), including published custom duplicates
    # that have no lesson sub-order. Keep prompts/images/phrases intact.
    lesson_numbers = {int(section.split("-", 1)[0]) for section in payloads}
    cleanup_urls.update(_clear_legacy_vocabulary_for_lessons(db, lesson_numbers))
    _remove_unreferenced_audio(db, cleanup_urls)
    return {"mode": mode, "published": published, **totals}


def build_vocabulary_import_template() -> bytes:
    """Build the admin's self-documenting XLSX import template."""
    workbook = openpyxl.Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions_rows = [
        ["Canonical vocabulary + quiz import template"],
        ["Workflow", "Upload this workbook first, then import a ZIP of audio files."],
        ["Identity", "Word Key is the stable identifier. Audio filenames must use the exact Word Key stem."],
        ["Rows", "Each Word Key must have exactly three rows: Round 1, Round 2 and Round 3."],
        ["Replace behavior", "The imported rows replace the canonical quiz bank for each Section. Existing audio is preserved by Word Key, and legacy scene/story vocabulary for the affected lesson is cleared."],
        ["Required question data", "Do not add Book source, Source Type, Tier, Skill Label, context or page-reference columns."],
        ["Round 1", "Meaning multiple choice: basic_meaning_mcq / click"],
        ["Round 2", "Pinyin typing: character_to_pinyin_typing / free_text"],
        ["Round 3", "Context cloze multiple choice: context_cloze_mcq / click"],
        ["Audio example", "C5-5-1-I1-W001.mp3"],
    ]
    for row in instructions_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 120
    instructions.freeze_panes = "A2"
    instructions["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    instructions.merge_cells("A1:B1")

    questions = workbook.create_sheet("Questions")
    questions.append(TEMPLATE_COLUMNS)
    sample_rows = [
        ["Q-DEMO-001", "C5-5-1-I1-W001", "5", "5-1", "1", "錢包", "qiánbāo", "N", "wallet; purse", "1", "basic_meaning_mcq", "mcq", "Choose the correct English meaning of 錢包.", "wallet; purse", "kitchen", "music", "chair", "A", "wallet; purse", "wallet; purse"],
        ["Q-DEMO-002", "C5-5-1-I1-W001", "5", "5-1", "1", "錢包", "qiánbāo", "N", "wallet; purse", "2", "character_to_pinyin_typing", "free_text", "Type the pinyin for 錢包.", "", "", "", "", "", "qiánbāo", "qiánbāo | qian2bao1"],
        ["Q-DEMO-003", "C5-5-1-I1-W001", "5", "5-1", "1", "錢包", "qiánbāo", "N", "wallet; purse", "3", "context_cloze_mcq", "mcq", "Choose the correct word in the sentence: 我找不到____。", "錢包", "鑰匙", "手機", "書", "A", "錢包", "錢包"],
    ]
    for row in sample_rows:
        questions.append(row)
    questions.freeze_panes = "A2"
    questions.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(TEMPLATE_COLUMNS))}{len(sample_rows) + 1}"
    for cell in questions[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="365C6B")
    for index, header in enumerate(TEMPLATE_COLUMNS, start=1):
        questions.column_dimensions[openpyxl.utils.get_column_letter(index)].width = max(14, min(42, len(header) + 4))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
