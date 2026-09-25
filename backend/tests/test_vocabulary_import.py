"""Admin quiz-vocabulary CSV/XLSX import: preview (read-only) and confirm (write)."""
import csv
import io
import zipfile

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from psycopg.types.json import Jsonb

import security.auth as auth
from db import connect_db
from routers import admin
from services.vocabulary_import import (
    _replace_assessment,
    apply_vocabulary_import,
    build_vocabulary_import_template,
    parse_csv_rows,
    parse_uploaded_rows,
    parse_xlsx_rows,
    preview_vocabulary_import,
    validate_import_rows,
)
from services.vocabulary_audio_import import (
    apply_vocabulary_audio_import,
    build_vocabulary_audio_sample,
    preview_vocabulary_audio_import,
)
import services.media as media_service
import services.vocabulary_audio_import as vocabulary_audio_service
from scripts.import_question_bank_workbook import build_payloads

COLUMNS = [
    "Question ID", "Word Key", "Source Type", "Chapter", "Section", "Item",
    "Traditional Chinese", "Pinyin", "POS", "English Meaning", "Round",
    "Question Type", "Input Mode", "Prompt", "Option A", "Option B",
    "Option C", "Option D", "Correct Option", "Correct Answer", "Accepted Answers",
]


def word_rows(*, word_key="C99-1-I1-W001", chinese="錢包", pinyin="qiánbāo", meaning="wallet", section="99-1", suffix=""):
    common = {
        "Word Key": word_key, "Source Type": "Vocabulary", "Chapter": section.split("-")[0],
        "Section": section, "Item": "1", "Traditional Chinese": chinese, "Pinyin": pinyin,
        "POS": "N", "English Meaning": meaning,
    }
    return [
        {**common, "Question ID": f"Q{word_key}1{suffix}", "Round": "Round 1",
         "Question Type": "basic_meaning_mcq", "Input Mode": "click",
         "Prompt": f"Choose the correct English meaning of {chinese}", "Option A": "kitchen",
         "Option B": "the front; the front side; ahead; in front", "Option C": meaning, "Option D": "chair",
         "Correct Option": "C", "Correct Answer": meaning, "Accepted Answers": meaning},
        {**common, "Question ID": f"Q{word_key}2{suffix}", "Round": "Round 2",
         "Question Type": "character_to_pinyin_typing", "Input Mode": "free_text",
         "Prompt": f"Type the pinyin for {chinese}.", "Option A": "", "Option B": "", "Option C": "", "Option D": "",
         "Correct Option": "", "Correct Answer": pinyin, "Accepted Answers": f"{pinyin} | {pinyin}1"},
        {**common, "Question ID": f"Q{word_key}3{suffix}", "Round": "Round 3",
         "Question Type": "context_cloze_mcq", "Input Mode": "click",
         "Prompt": f"我的___不見了。", "Option A": chinese, "Option B": "書", "Option C": "門", "Option D": "床",
         "Correct Option": "A", "Correct Answer": chinese, "Accepted Answers": chinese},
    ]


def csv_bytes(rows: list[dict[str, str]], columns: list[str] = COLUMNS) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def xlsx_bytes(rows: list[dict[str, str]], *, sheet_name: str = "Questions", columns: list[str] = COLUMNS) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_xlsx_rows_matches_csv_rows():
    rows = word_rows()
    assert parse_xlsx_rows(xlsx_bytes(rows)) == rows


def test_parse_xlsx_rows_finds_a_sheet_named_questions_even_if_not_first():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Cover"
    sheet = workbook.create_sheet("Questions")
    sheet.append(COLUMNS)
    for row in word_rows():
        sheet.append([row.get(column, "") for column in COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert parse_xlsx_rows(buffer.getvalue()) == word_rows()


def test_parse_xlsx_rows_accepts_export_ready_numeric_rounds_and_mcq_mode():
    rows = word_rows()
    for row in rows:
        row["Round"] = row["Round"].split()[-1]
        if row["Input Mode"] == "click":
            row["Input Mode"] = "mcq"
    parsed = parse_xlsx_rows(xlsx_bytes(rows, sheet_name="Quiz Questions"))
    assert parsed == word_rows()


def test_import_payload_contains_only_canonical_question_data():
    rows = word_rows()
    for content, filename in ((csv_bytes(rows), "bank.csv"), (xlsx_bytes(rows), "bank.xlsx")):
        parsed = parse_uploaded_rows(filename, content)
        assert validate_import_rows(parsed) == []
        payload = build_payloads(parsed)["99-1"]
        assert all(set(question) == {
            "questionId", "wordId", "targetWord", "pinyin", "pos", "simpleEnglishMeaning",
            "level", "difficultyWeight", "questionType", "answerFormat", "prompt", "options",
            "correctAnswer", "acceptedAnswers", "explanation", "sourceQuestionId", "sourceType", "round",
        } for question in payload)


def test_parse_xlsx_rows_reports_missing_columns():
    workbook = openpyxl.Workbook()
    workbook.active.append(["Question ID", "Word Key"])
    workbook.active.append(["Q1", "W1"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ValueError, match="missing required columns"):
        parse_xlsx_rows(buffer.getvalue())


def test_parse_uploaded_rows_dispatches_by_extension():
    rows = word_rows()
    assert parse_uploaded_rows("bank.csv", csv_bytes(rows)) == rows
    assert parse_uploaded_rows("bank.xlsx", xlsx_bytes(rows)) == rows
    assert parse_uploaded_rows("", csv_bytes(rows)) == rows  # no extension falls back to CSV


def test_parse_uploaded_rows_rejects_legacy_xls():
    with pytest.raises(ValueError, match="legacy .xls format is not supported"):
        parse_uploaded_rows("bank.xls", b"anything")


def test_preview_and_confirm_from_an_xlsx_file(admin_client):
    story = {
        "id": "vocab-import-xlsx-story-98-1", "title": "XLSX import test", "frames": [],
        "published": True, "lessonNumber": 98, "lessonSubOrder": 1,
    }
    assert admin_client.post("/api/custom-stories", json=story).status_code == 200
    content = xlsx_bytes(word_rows(word_key="C98-1-I1-W001", section="98-1"))
    with connect_db() as db:
        preview = preview_vocabulary_import(db, content, filename="bank.xlsx")
    assert preview["rowIssues"] == []
    assert preview["sections"][0]["newWords"] == 1
    with connect_db() as db:
        result = apply_vocabulary_import(db, content, filename="bank.xlsx")
    assert result["published"][0]["storyId"] == "vocab-import-xlsx-story-98-1"


def test_parse_csv_rows_reports_missing_columns():
    bad = b"Question ID,Word Key\nQ1,W1\n"
    with pytest.raises(ValueError, match="missing required columns"):
        parse_csv_rows(bad)


def test_parse_csv_rows_accepts_template_without_source_type():
    columns = [column for column in COLUMNS if column != "Source Type"]
    parsed = parse_csv_rows(csv_bytes(word_rows(), columns=columns))
    assert validate_import_rows(parsed) == []


def test_validate_import_rows_accepts_a_well_formed_word():
    rows = word_rows()
    assert validate_import_rows(rows) == []


def test_validate_import_rows_flags_wrong_question_type_for_round():
    rows = word_rows()
    rows[1]["Question Type"] = "context_cloze_mcq"  # should be character_to_pinyin_typing
    issues = validate_import_rows(rows)
    assert any("expected character_to_pinyin_typing" in issue for issue in issues)


def test_validate_import_rows_flags_missing_round():
    rows = word_rows()[:2]  # drop Round 3
    issues = validate_import_rows(rows)
    assert any("expected exactly one row for each round" in issue for issue in issues)


def test_replace_assessment_removes_words_outside_the_lesson_file_and_preserves_audio():
    existing = [
        {"wordId": "W001", "targetWord": "old", "audioUrl": "/uploads/audio/w001.mp3"},
        {"wordId": "W001", "targetWord": "old", "audioUrl": "/uploads/audio/w001.mp3"},
        {"wordId": "W001", "targetWord": "old", "audioUrl": "/uploads/audio/w001.mp3"},
        {"wordId": "W002", "targetWord": "remove me", "audioUrl": "/uploads/audio/w002.mp3"},
    ]
    incoming = [{"wordId": "W001", "targetWord": "new"}]
    replacement, new_words, removed_words, missing_audio = _replace_assessment(existing, incoming)
    assert replacement == [{"wordId": "W001", "targetWord": "new", "audioUrl": "/uploads/audio/w001.mp3"}]
    assert new_words == set()
    assert removed_words == {"W002"}
    assert missing_audio == set()


def test_preview_reports_row_issues_without_touching_the_database(client):
    rows = word_rows()
    rows[1]["Question Type"] = "context_cloze_mcq"
    with connect_db() as db:
        report = preview_vocabulary_import(db, csv_bytes(rows))
    assert report["sections"] == []
    assert any("expected character_to_pinyin_typing" in issue for issue in report["rowIssues"])


def test_preview_and_confirm_against_a_real_story(admin_client):
    story = {
        "id": "vocab-import-story-99-1", "title": "Import test story", "frames": [],
        "published": True, "lessonNumber": 99, "lessonSubOrder": 1,
    }
    assert admin_client.post("/api/custom-stories", json=story).status_code == 200

    rows = word_rows()
    content = csv_bytes(rows)
    with connect_db() as db:
        preview = preview_vocabulary_import(db, content)
    assert preview["rowIssues"] == []
    assert len(preview["sections"]) == 1
    section = preview["sections"][0]
    assert section["found"] is True
    assert section["storyId"] == "vocab-import-story-99-1"
    assert section["newWords"] == 1
    assert section["updatedWords"] == 0
    assert section["issues"] == []

    with connect_db() as db:
        result = apply_vocabulary_import(db, content)
    assert result["published"] == [{
        "section": "99-1", "storyId": "vocab-import-story-99-1", "storyTitle": "Import test story",
        "questionCount": 3, "newWords": 1, "updatedWords": 0, "removedWords": 0,
        "preservedAudio": 0, "missingAudio": 1,
    }]
    assert result["mode"] == "replace_lesson"

    saved = next(s for s in admin_client.get("/api/custom-stories").json() if s["id"] == "vocab-import-story-99-1")
    assert len(saved["vocabAssessment"]) == 3
    assert {q["wordId"] for q in saved["vocabAssessment"]} == {"C99-1-I1-W001"}

    # Re-importing the same word replaces its three rounds rather than duplicating it.
    updated_rows = word_rows(meaning="coin purse")
    with connect_db() as db:
        preview_again = preview_vocabulary_import(db, csv_bytes(updated_rows))
    assert preview_again["sections"][0]["newWords"] == 0
    assert preview_again["sections"][0]["updatedWords"] == 1

    with connect_db() as db:
        apply_vocabulary_import(db, csv_bytes(updated_rows))
    saved_again = next(s for s in admin_client.get("/api/custom-stories").json() if s["id"] == "vocab-import-story-99-1")
    assert len(saved_again["vocabAssessment"]) == 3
    assert saved_again["vocabAssessment"][0]["simpleEnglishMeaning"] == "coin purse"
    assert all("sourceType" not in question for question in saved_again["vocabAssessment"])


def test_replace_lesson_removes_stale_words_and_preserves_story_content(admin_client, tmp_path, monkeypatch):
    story = {
        "id": "vocab-replace-story-99-2", "title": "Replace test story", "frames": [],
        "published": True, "lessonNumber": 99, "lessonSubOrder": 2,
    }
    assert admin_client.post("/api/custom-stories", json=story).status_code == 200
    upload_root = tmp_path / "uploads"
    old_audio = upload_root / "audio" / "old.mp3"
    old_audio.parent.mkdir(parents=True)
    old_audio.write_bytes(b"old")
    monkeypatch.setattr(media_service, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(media_service, "AUDIO_UPLOAD_DIR", str(upload_root / "audio"))
    original_frames = [{"imageUrl": "frame.png", "prompt": "Keep this frame"}]
    original_story_vocabulary = {"easy": {"vocabulary": "keep this story vocab"}}
    old_assessment = [
        {"wordId": "OLD-W001", "audioUrl": "/uploads/audio/old.mp3"},
        {"wordId": "OLD-W001", "audioUrl": "/uploads/audio/old.mp3"},
        {"wordId": "OLD-W001", "audioUrl": "/uploads/audio/old.mp3"},
    ]
    with connect_db() as db:
        db.execute(
            "UPDATE custom_stories SET frames = %s::jsonb, story_vocabulary = %s::jsonb, vocab_assessment = %s::jsonb WHERE id = %s",
            (Jsonb(original_frames), Jsonb(original_story_vocabulary), Jsonb(old_assessment), story["id"]),
        )
        result = apply_vocabulary_import(db, csv_bytes(word_rows(word_key="NEW-W001", section="99-2")))

    assert result["removedWords"] == 1
    assert result["missingAudio"] == 1
    assert not old_audio.exists()
    saved = next(item for item in admin_client.get("/api/custom-stories").json() if item["id"] == story["id"])
    assert {question["wordId"] for question in saved["vocabAssessment"]} == {"NEW-W001"}
    assert saved["frames"] == original_frames
    assert saved["storyVocabulary"] == original_story_vocabulary


def test_preview_reports_when_no_story_matches_the_section(client):
    rows = word_rows(section="97-9")
    with connect_db() as db:
        report = preview_vocabulary_import(db, csv_bytes(rows))
    assert report["sections"][0]["found"] is False
    assert "No existing story" in report["sections"][0]["error"]


def test_confirm_raises_when_no_story_matches_the_section(client):
    rows = word_rows(section="97-9")
    with connect_db() as db, pytest.raises(LookupError):
        apply_vocabulary_import(db, csv_bytes(rows))


def audio_zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for filename, content in files.items():
            archive.writestr(filename, content)
    return buffer.getvalue()


def test_vocabulary_audio_import_matches_word_keys_and_updates_all_rounds(admin_client, tmp_path, monkeypatch):
    story = {
        "id": "vocab-audio-import-story-99-1", "title": "Audio import story", "frames": [],
        "published": True, "lessonNumber": 99, "lessonSubOrder": 1,
    }
    assert admin_client.post("/api/custom-stories", json=story).status_code == 200
    rows = word_rows(word_key="C99-1-I1-W001")
    with connect_db() as db:
        apply_vocabulary_import(db, csv_bytes(rows))

    upload_root = tmp_path / "uploads"
    monkeypatch.setattr(media_service, "UPLOAD_DIR", str(upload_root))
    monkeypatch.setattr(media_service, "AUDIO_UPLOAD_DIR", str(upload_root / "audio"))
    content = audio_zip({"C99-1-I1-W001.mp3": b"fake-mp3", "unknown.mp3": b"ignored"})
    with connect_db() as db:
        preview = preview_vocabulary_audio_import(db, content)
    assert preview["files"] == 2
    assert [item["wordKey"] for item in preview["matched"]] == ["C99-1-I1-W001"]
    assert preview["unmatched"] == ["unknown.mp3"]

    with connect_db() as db:
        result = apply_vocabulary_audio_import(db, content)
    assert result["updated"] == 1
    saved = next(s for s in admin_client.get("/api/custom-stories").json() if s["id"] == story["id"])
    audio_urls = {question["audioUrl"] for question in saved["vocabAssessment"]}
    assert len(audio_urls) == 1
    first_url = next(iter(audio_urls))
    assert first_url.startswith("/uploads/audio/vocab-C99-1-I1-W001-")
    first_path = (upload_root / "audio").joinpath(first_url.rsplit("/", 1)[-1])
    assert first_path.is_file()

    with connect_db() as db:
        result_again = apply_vocabulary_audio_import(db, audio_zip({"C99-1-I1-W001.mp3": b"replacement-mp3"}))
    assert result_again["updated"] == 1
    saved_again = next(s for s in admin_client.get("/api/custom-stories").json() if s["id"] == story["id"])
    second_urls = {question["audioUrl"] for question in saved_again["vocabAssessment"]}
    assert len(second_urls) == 1
    second_url = next(iter(second_urls))
    assert second_url != first_url
    assert not first_path.exists()
    assert (upload_root / "audio").joinpath(second_url.rsplit("/", 1)[-1]).is_file()


def test_audio_sample_preview_matches_the_template_word_key(monkeypatch):
    monkeypatch.setattr(
        vocabulary_audio_service,
        "_word_locations",
        lambda _db: {
            "C5-5-1-I1-W001": [{"storyId": "sample-story", "storyTitle": "Sample lesson", "questionIndex": 0}]
        },
    )
    preview = preview_vocabulary_audio_import(object(), build_vocabulary_audio_sample())
    assert preview["files"] == 1
    assert [match["wordKey"] for match in preview["matched"]] == ["C5-5-1-I1-W001"]
    assert preview["unmatched"] == []


def test_audio_word_key_suffix_and_duplicates_are_not_accepted(admin_client, tmp_path, monkeypatch):
    story = {
        "id": "vocab-audio-validation-story-99-2", "title": "Audio validation story", "frames": [],
        "published": True, "lessonNumber": 99, "lessonSubOrder": 2,
    }
    assert admin_client.post("/api/custom-stories", json=story).status_code == 200
    with connect_db() as db:
        db.execute(
            "UPDATE custom_stories SET vocab_assessment = %s::jsonb WHERE id = %s",
            (Jsonb([
                {"wordId": "C99-2-I1-W001"},
                {"wordId": "C99-2-I1-W001"},
                {"wordId": "C99-2-I1-W001"},
            ]), story["id"]),
        )
        suffix_preview = preview_vocabulary_audio_import(
            db, audio_zip({"C99-2-I1-W001_r1.mp3": b"wrong-key"})
        )
        assert suffix_preview["matched"] == []
        assert suffix_preview["unmatched"] == ["C99-2-I1-W001_r1.mp3"]

        duplicate = audio_zip({"C99-2-I1-W001.mp3": b"one", "C99-2-I1-W001.wav": b"two"})
        with pytest.raises(ValueError, match="duplicate audio files"):
            apply_vocabulary_audio_import(db, duplicate)


@pytest.fixture()
def import_endpoint_api(monkeypatch):
    app = FastAPI()
    app.include_router(admin.router)
    with TestClient(app) as test_client:
        yield test_client


def test_import_endpoints_are_admin_only(import_endpoint_api, monkeypatch):
    test_client = import_endpoint_api
    monkeypatch.setattr(admin, "preview_vocabulary_import", lambda db, content, filename="", mode="": {
        "mode": mode, "rows": 0, "rowIssues": [], "sections": [], "newWords": 0,
        "updatedWords": 0, "removedWords": 0, "preservedAudio": 0, "missingAudio": 0,
    })
    files = {"file": ("bank.csv", b"Question ID\n", "text/csv"), "mode": (None, "replace_lesson")}

    response = test_client.post("/api/admin/vocabulary-import/preview", files=files)
    assert response.status_code == 401

    test_client.cookies.set(auth.ROLE_COOKIE_NAMES["teacher"], auth.issue_token("teacher", "teacher-1"))
    assert test_client.post("/api/admin/vocabulary-import/preview", files=files).status_code in (401, 403)

    test_client.cookies.clear()
    test_client.cookies.set(auth.ROLE_COOKIE_NAMES["admin"], auth.issue_token("admin", "admin"))
    response = test_client.post("/api/admin/vocabulary-import/preview", files=files)
    assert response.status_code == 200
    assert response.json() == {
        "mode": "replace_lesson", "rows": 0, "rowIssues": [], "sections": [], "newWords": 0,
        "updatedWords": 0, "removedWords": 0, "preservedAudio": 0, "missingAudio": 0,
    }


def test_preview_endpoint_returns_422_on_bad_file(import_endpoint_api):
    test_client = import_endpoint_api
    test_client.cookies.set(auth.ROLE_COOKIE_NAMES["admin"], auth.issue_token("admin", "admin"))
    files = {"file": ("bank.csv", b"not,a,valid,header\n", "text/csv"), "mode": (None, "replace_lesson")}
    response = test_client.post("/api/admin/vocabulary-import/preview", files=files)
    assert response.status_code == 422
    assert "missing required columns" in response.json()["detail"]


def test_template_has_instructions_and_questions_sheets_without_retired_source_column():
    workbook = openpyxl.load_workbook(io.BytesIO(build_vocabulary_import_template()), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Instructions", "Questions"]
    headers = [cell.value for cell in next(workbook["Questions"].iter_rows(max_row=1))]
    assert "Word Key" in headers
    assert "Source Type" not in headers
    assert "Book source" not in headers
    workbook.close()


def test_template_endpoint_is_admin_only(import_endpoint_api):
    test_client = import_endpoint_api
    assert test_client.get("/api/admin/vocabulary-import/template").status_code == 401
    test_client.cookies.set(auth.ROLE_COOKIE_NAMES["admin"], auth.issue_token("admin", "admin"))
    response = test_client.get("/api/admin/vocabulary-import/template")
    assert response.status_code == 200
    assert response.headers["content-disposition"].endswith('vocabulary-import-template.xlsx"')
    workbook = openpyxl.load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Instructions", "Questions"]
    workbook.close()


def test_audio_template_contains_only_the_canonical_word_key_filename():
    with zipfile.ZipFile(io.BytesIO(build_vocabulary_audio_sample())) as archive:
        assert archive.namelist() == ["C5-5-1-I1-W001.mp3"]
        assert archive.read("C5-5-1-I1-W001.mp3").startswith(b"Mapping-only sample")


def test_audio_template_endpoint_is_admin_only(import_endpoint_api):
    test_client = import_endpoint_api
    assert test_client.get("/api/admin/vocabulary-audio-import/template").status_code == 401
    test_client.cookies.set(auth.ROLE_COOKIE_NAMES["admin"], auth.issue_token("admin", "admin"))
    response = test_client.get("/api/admin/vocabulary-audio-import/template")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/zip")
    assert response.headers["content-disposition"].endswith('vocabulary-audio-sample.zip"')
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.namelist() == ["C5-5-1-I1-W001.mp3"]
