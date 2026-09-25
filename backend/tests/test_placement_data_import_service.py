from io import BytesIO

import openpyxl

from services import placement_data_import_service as service


def _blueprint():
    return {
        "revision": 4,
        "questions": [
            {
                "questionId": "Q0001",
                "sourceStoryId": "story-1",
                "sourceWordId": "word-1",
                "position": 1,
                "questionType": "basic_meaning_mcq",
                "round": 1,
                "tier": "tier1",
                "correctAnswer": "book",
                "acceptedAnswers": ["book"],
                "options": ["book", "chair"],
            },
            {
                "questionId": "Q0002",
                "sourceStoryId": "story-2",
                "sourceWordId": "word-2",
                "position": 2,
                "questionType": "context_cloze_mcq",
                "round": 3,
                "tier": "tier3",
                "correctAnswer": "chair",
                "acceptedAnswers": ["chair"],
                "options": ["chair", "book"],
            },
        ],
    }


def test_sample_workbook_uses_exact_response_headers(monkeypatch):
    monkeypatch.setattr(service.workbook_import, "EXPECTED_STUDENT_IDS", ("SIM001", "SIM002"))
    monkeypatch.setattr(service.workbook_import, "load_active_blueprint", lambda _db: _blueprint())

    content = service.build_sample_workbook(object())
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook["Responses"]

    assert [cell.value for cell in next(sheet.iter_rows())] == list(service.workbook_import.EXPECTED_HEADERS)
    assert sum(1 for _ in sheet.iter_rows()) == 5
    assert workbook.sheetnames == ["Responses", "Instructions"]
    workbook.close()


def test_preview_import_returns_a_safe_validation_payload(monkeypatch):
    plan = {
        "attempts": [{"id": "session-1"}],
        "summary": {
            "student_count": 1,
            "question_count": 2,
            "response_count": 2,
            "correct_count": 1,
            "incorrect_count": 1,
            "correct_by_mode": {"tier1": 1},
            "source_aliases": {},
        },
    }
    monkeypatch.setattr(service, "_build_plan", lambda *_args: (plan, {"skip_sessions": set()}))

    result = service.preview_import(object(), b"workbook", "responses.xlsx")

    assert result["valid"] is True
    assert result["studentCount"] == 1
    assert result["responseCount"] == 2
    assert result["newSessions"] == 1
